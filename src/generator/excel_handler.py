import openpyxl


def leer_excel(ruta_excel):
    """Lee el Excel y retorna (encabezados, lista de diccionarios)."""
    wb = openpyxl.load_workbook(ruta_excel, read_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(filas) < 2:
        raise ValueError(
            "El archivo Excel debe tener al menos una fila de encabezados y una fila de datos."
        )

    encabezados = [str(c).strip() for c in filas[0]]
    datos = []
    for fila in filas[1:]:
        registro = {}
        for j, valor in enumerate(fila):
            if j < len(encabezados):
                registro[encabezados[j]] = valor if valor is not None else ""
        datos.append(registro)
    return encabezados, datos


def validar_columnas(columnas_excel, etiquetas_pptx):
    """Valida que las columnas del Excel coincidan con las etiquetas del PPTX.
    Retorna (es_valido, faltantes, sobrantes)."""
    set_excel = set(columnas_excel)
    set_pptx = set(etiquetas_pptx)
    faltantes = sorted(set_pptx - set_excel)
    sobrantes = sorted(set_excel - set_pptx)
    es_valido = len(faltantes) == 0
    return es_valido, faltantes, sobrantes


def generar_plantilla_excel(etiquetas, ruta_salida):
    """Genera un Excel con los encabezados de las etiquetas detectadas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Participantes"
    for col, etiqueta in enumerate(etiquetas, start=1):
        celda = ws.cell(row=1, column=col, value=etiqueta)
        celda.font = openpyxl.styles.Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 25
    wb.save(ruta_salida)
    wb.close()
