import copy
import io
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def extraer_etiquetas(ruta_plantilla):
    """Extrae todas las etiquetas {{...}} de la plantilla PPTX."""
    prs = Presentation(ruta_plantilla)
    etiquetas = set()
    patron = re.compile(r"\{\{(.+?)\}\}")

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        encontradas = patron.findall(run.text)
                        etiquetas.update(encontradas)
    return sorted(etiquetas)


def duplicar_slide(prs, slide):
    """Duplica una diapositiva copiando formas y reconstruyendo
    las relaciones de las imágenes para que no se pierdan."""
    layout = slide.slide_layout
    nueva = prs.slides.add_slide(layout)

    for shape in list(nueva.shapes):
        el = shape._element
        el.getparent().remove(el)

    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image_stream = io.BytesIO(shape.image.blob)
            nueva.shapes.add_picture(
                image_stream, shape.left, shape.top, shape.width, shape.height
            )
        else:
            nuevo_el = copy.deepcopy(shape._element)
            nueva.shapes._spTree.insert_element_before(nuevo_el, "p:extLst")

    return nueva


def reemplazar_etiquetas(slide, datos):
    """Reemplaza todas las etiquetas {{...}} del slide con los datos."""
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    for clave, valor in datos.items():
                        run.text = run.text.replace(f"{{{{{clave}}}}}", str(valor))


def generar_certificados(lista_participantes, ruta_plantilla, archivo_salida):
    """Genera el PPTX con un slide por cada participante."""
    prs = Presentation(ruta_plantilla)
    plantilla = prs.slides[0]

    for persona in lista_participantes:
        slide = duplicar_slide(prs, plantilla)
        reemplazar_etiquetas(slide, persona)

    id_slide = prs.slides._sldIdLst[0]
    prs.slides._sldIdLst.remove(id_slide)

    prs.save(archivo_salida)
    return len(lista_participantes)
